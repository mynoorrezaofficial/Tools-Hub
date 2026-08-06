import os
from datetime import datetime


PDF_EDITABLE_FIELDS = {
    "Title": "/Title",
    "Author": "/Author",
    "Subject": "/Subject",
    "Keywords": "/Keywords",
    "Creator": "/Creator",
    "Producer": "/Producer",
}

JPEG_EDITABLE_FIELDS = {
    "ImageDescription": ("0th", 270),
    "Artist": ("0th", 315),
    "Copyright": ("0th", 33432),
    "Software": ("0th", 305),
    "DateTime": ("0th", 306),
}

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx"}
AUDIO_EXTENSIONS = {".mp3", ".m4a", ".aac", ".flac", ".ogg", ".wav", ".wma"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v"}


def _safe_text(value):
    if value is None:
        return None
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="ignore").rstrip("\x00")
    if isinstance(value, (list, tuple)):
        return ", ".join([_safe_text(v) for v in value if _safe_text(v) not in (None, "")])
    if isinstance(value, dict):
        return {k: _safe_text(v) for k, v in value.items()}
    return str(value)


def _format_timestamp(ts):
    try:
        return datetime.fromtimestamp(ts).isoformat(sep=" ", timespec="seconds")
    except Exception:
        return None


def _guess_true_type(file_path):
    try:
        import filetype

        kind = filetype.guess(file_path)
        if kind:
            label = f"{kind.extension.upper()} ({kind.mime})"
            return {"mime": kind.mime, "extension": kind.extension, "label": label}
    except Exception:
        pass

    ext = os.path.splitext(file_path)[1].lower().strip('.')
    return {"mime": None, "extension": ext or None, "label": ext.upper() if ext else "Unknown"}


def _file_info_section(file_path):
    stat = os.stat(file_path)
    guessed = _guess_true_type(file_path)
    return {
        "File Name": os.path.basename(file_path),
        "File Size": f"{stat.st_size} bytes",
        "Extension": os.path.splitext(file_path)[1].lower(),
        "True File Type": guessed["label"],
        "MIME Type": guessed["mime"],
        "Created": _format_timestamp(stat.st_ctime),
        "Modified": _format_timestamp(stat.st_mtime),
    }


def _collect_hachoir_map(file_path):
    from hachoir.parser import createParser
    from hachoir.metadata import extractMetadata

    parser = createParser(file_path)
    if not parser:
        return {}

    try:
        with parser:
            metadata = extractMetadata(parser)
    except Exception:
        return {}

    if not metadata or not hasattr(metadata, 'metadata') or not metadata.metadata:
        return {}

    collected = {}
    try:
        keys = metadata.metadata.keys()
    except Exception:
        keys = []

    for key in keys:
        try:
            value = metadata.get(key)
            if value is not None:
                collected[str(key)] = _safe_text(value)
        except Exception:
            pass

    return collected


def _flatten_sections(sections):
    flat = {}
    for section_name, section_data in sections.items():
        if not isinstance(section_data, dict):
            flat[section_name] = _safe_text(section_data)
            continue
        for key, value in section_data.items():
            if value is None or value == "":
                continue
            flat[key] = _safe_text(value)
    return flat


def _build_response(file_type, sections):
    return {
        "file_type": file_type,
        "sections": sections,
        "metadata": _flatten_sections(sections),
    }


def extract_metadata(file_path):
    """Extract metadata from a file."""
    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == ".pdf":
            return _extract_pdf_metadata(file_path)
        if ext in (".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp"):
            return _extract_image_metadata(file_path)
        if ext in (".docx", ".xlsx", ".pptx"):
            return _extract_office_metadata(file_path)
        if ext in (".mp3", ".m4a", ".aac", ".flac", ".ogg", ".wav", ".wma"):
            return _extract_audio_metadata(file_path)
        if ext in (".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v"):
            return _extract_video_metadata(file_path)
        return _extract_generic_metadata(file_path)
    except Exception as e:
        return {"error": str(e)}


def update_pdf_metadata(file_path, output_path, metadata_updates=None, delete_all=False):
    """Replace or delete PDF metadata using pikepdf."""
    import pikepdf

    metadata_updates = metadata_updates or {}

    try:
        with pikepdf.open(file_path) as pdf:
            def clear_docinfo():
                for key in list(pdf.docinfo.keys()):
                    try:
                        del pdf.docinfo[key]
                    except Exception:
                        pass

            if delete_all:
                clear_docinfo()
                try:
                    if "/Metadata" in pdf.Root:
                        del pdf.Root.Metadata
                except Exception:
                    pass
            else:
                clear_docinfo()
                for key, value in metadata_updates.items():
                    if value is None:
                        continue
                    text = str(value).strip()
                    if text == "":
                        continue
                    try:
                        pdf.docinfo[f"/{key}"] = text
                    except Exception:
                        from pikepdf import Name
                        pdf.docinfo[Name(f"/{key}")] = text

            pdf.save(output_path)

        return True, "PDF metadata updated successfully."
    except Exception as e:
        return False, str(e)


def update_jpeg_metadata(file_path, output_path, metadata_updates=None, delete_all=False):
    """Replace or delete JPEG EXIF metadata using piexif."""
    from PIL import Image
    import piexif

    metadata_updates = metadata_updates or {}

    try:
        image = Image.open(file_path)
        existing_exif = image.info.get("exif")
        if existing_exif:
            exif_dict = piexif.load(existing_exif)
        else:
            exif_dict = {"0th": {}, "Exif": {}, "GPS": {}, "1st": {}, "thumbnail": None}

        if delete_all:
            exif_dict = {"0th": {}, "Exif": {}, "GPS": {}, "1st": {}, "thumbnail": None}
        else:
            for field, (ifd, tag) in JPEG_EDITABLE_FIELDS.items():
                value = metadata_updates.get(field)
                if value is None:
                    continue
                text = str(value).strip()
                if not text:
                    continue
                exif_dict[ifd][tag] = text.encode("utf-8", errors="ignore")

        exif_bytes = piexif.dump(exif_dict)
        image.save(output_path, format="JPEG", exif=exif_bytes, quality=95)
        return True, "JPEG metadata updated successfully."
    except Exception as e:
        return False, str(e)


def _extract_pdf_metadata(file_path):
    """Extract metadata from PDF using pikepdf."""
    import pikepdf

    with pikepdf.open(file_path) as pdf:
        meta = pdf.docinfo or {}
        info = {
            "file_type": "PDF",
            "pages": len(pdf.pages),
            "metadata": {}
        }

        mapping = {
            "/Title": "Title",
            "/Author": "Author",
            "/Subject": "Subject",
            "/Creator": "Creator",
            "/Producer": "Producer",
            "/CreationDate": "Creation Date",
            "/ModDate": "Modification Date",
            "/Keywords": "Keywords",
        }

        for raw_key, label in mapping.items():
            val = meta.get(raw_key)
            if val is not None:
                try:
                    info["metadata"][label] = str(val)
                except Exception:
                    info["metadata"][label] = repr(val)

        sections = {
            "file_info": _file_info_section(file_path),
            "document": info["metadata"],
            "technical": {"Pages": len(pdf.pages)},
        }

        hachoir_map = _collect_hachoir_map(file_path)
        if hachoir_map:
            sections["technical"].update({k: v for k, v in hachoir_map.items() if k not in sections["document"]})

        if not sections["document"]:
            sections["document"]["Note"] = "No embedded metadata found in PDF."

        return _build_response("PDF", sections)


def _extract_image_metadata(file_path):
    from PIL import Image

    ext = os.path.splitext(file_path)[1].upper().strip('.')
    sections = {"file_info": _file_info_section(file_path)}

    try:
        image = Image.open(file_path)
        sections["image"] = {
            "Dimensions": f"{image.width} x {image.height}",
            "Mode": image.mode,
        }
    except Exception:
        sections["image"] = {}

    if ext in ("JPG", "JPEG"):
        import piexif

        try:
            exif_dict = piexif.load(file_path)
        except Exception:
            exif_dict = None

        if exif_dict:
            def gps_to_decimal(values, ref):
                try:
                    deg = values[0][0] / values[0][1]
                    minute = values[1][0] / values[1][1]
                    second = values[2][0] / values[2][1]
                    decimal = deg + minute / 60 + second / 3600
                    if ref in (b'S', b'W', 'S', 'W'):
                        decimal *= -1
                    return round(decimal, 6)
                except Exception:
                    return None

            zeroth = exif_dict.get("0th", {})
            exif_ifd = exif_dict.get("Exif", {})
            gps_ifd = exif_dict.get("GPS", {})

            camera_model = zeroth.get(piexif.ImageIFD.Model)
            software = zeroth.get(piexif.ImageIFD.Software)
            date_taken = exif_ifd.get(piexif.ExifIFD.DateTimeOriginal) or zeroth.get(piexif.ImageIFD.DateTime)
            exposure = exif_ifd.get(piexif.ExifIFD.ExposureTime)
            f_number = exif_ifd.get(piexif.ExifIFD.FNumber)
            iso = exif_ifd.get(piexif.ExifIFD.ISOSpeedRatings)
            lens = exif_ifd.get(piexif.ExifIFD.LensModel)

            gps_lat = gps_to_decimal(gps_ifd.get(piexif.GPSIFD.GPSLatitude), gps_ifd.get(piexif.GPSIFD.GPSLatitudeRef)) if gps_ifd.get(piexif.GPSIFD.GPSLatitude) else None
            gps_lon = gps_to_decimal(gps_ifd.get(piexif.GPSIFD.GPSLongitude), gps_ifd.get(piexif.GPSIFD.GPSLongitudeRef)) if gps_ifd.get(piexif.GPSIFD.GPSLongitude) else None

            sections["image"].update({
                "Camera Make": _safe_text(zeroth.get(piexif.ImageIFD.Make)),
                "Camera Model": _safe_text(camera_model),
                "Lens Model": _safe_text(lens),
                "Date/Time Taken": _safe_text(date_taken),
                "Exposure Time": _safe_text(exposure),
                "F Number": _safe_text(f_number),
                "ISO": _safe_text(iso),
                "Software Used": _safe_text(software),
                "GPS Latitude": gps_lat,
                "GPS Longitude": gps_lon,
            })

    hachoir_map = _collect_hachoir_map(file_path)
    if hachoir_map:
        technical = {}
        for key, value in hachoir_map.items():
            if any(term in key.lower() for term in ("date", "time", "model", "software", "creator", "gps", "location", "dpi", "compression", "color")):
                technical[key] = value
        if technical:
            sections["technical"] = technical

    if len(sections.get("image", {})) <= 2:
        sections["image"]["Note"] = "No EXIF metadata found."

    return _build_response(ext or "IMAGE", sections)


def _extract_office_metadata(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    sections = {"file_info": _file_info_section(file_path), "document": {}}

    try:
        if ext == ".docx":
            from docx import Document

            props = Document(file_path).core_properties
            sections["document"] = {
                "Author": _safe_text(props.author),
                "Last Modified By": _safe_text(props.last_modified_by),
                "Created": _safe_text(props.created),
                "Modified": _safe_text(props.modified),
                "Title": _safe_text(props.title),
                "Keywords": _safe_text(props.keywords),
                "Creator": _safe_text(props.author),
                "Software Used": "Microsoft Word / Word-compatible DOCX",
            }
        elif ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(file_path, read_only=True)
            props = wb.properties
            sections["document"] = {
                "Author": _safe_text(getattr(props, "creator", None)),
                "Last Modified By": _safe_text(getattr(props, "lastModifiedBy", None)),
                "Created": _safe_text(getattr(props, "created", None)),
                "Modified": _safe_text(getattr(props, "modified", None)),
                "Title": _safe_text(getattr(props, "title", None)),
                "Keywords": _safe_text(getattr(props, "keywords", None)),
                "Software Used": _safe_text(getattr(props, "application", None)),
            }
            wb.close()
        elif ext == ".pptx":
            from pptx import Presentation

            props = Presentation(file_path).core_properties
            sections["document"] = {
                "Author": _safe_text(props.author),
                "Last Modified By": _safe_text(props.last_modified_by),
                "Created": _safe_text(props.created),
                "Modified": _safe_text(props.modified),
                "Title": _safe_text(props.title),
                "Keywords": _safe_text(props.keywords),
                "Software Used": _safe_text(props.creator),
            }
    except Exception as e:
        sections["document"]["Note"] = str(e)

    hachoir_map = _collect_hachoir_map(file_path)
    if hachoir_map:
        sections["technical"] = {k: v for k, v in hachoir_map.items() if k not in sections["document"]}

    return _build_response(ext.upper().strip('.') or "DOCUMENT", sections)


def _extract_audio_metadata(file_path):
    from mutagen import File as MutagenFile

    ext = os.path.splitext(file_path)[1].lower()
    sections = {"file_info": _file_info_section(file_path), "audio": {}}

    try:
        audio = MutagenFile(file_path, easy=True)
        if audio:
            tags = audio.tags or {}
            sections["audio"] = {
                "Track Title": _safe_text((tags.get("title") or [None])[0]),
                "Artist": _safe_text((tags.get("artist") or [None])[0]),
                "Album": _safe_text((tags.get("album") or [None])[0]),
                "Genre": _safe_text((tags.get("genre") or [None])[0]),
                "Year": _safe_text((tags.get("date") or tags.get("year") or [None])[0]),
                "Duration": round(getattr(audio.info, "length", 0), 2) if getattr(audio, "info", None) else None,
            }
    except Exception as e:
        sections["audio"]["Note"] = str(e)

    hachoir_map = _collect_hachoir_map(file_path)
    if hachoir_map:
        technical = {}
        for key, value in hachoir_map.items():
            if any(term in key.lower() for term in ("duration", "bitrate", "sample", "channel", "album", "artist", "title", "genre", "year", "track")):
                technical[key] = value
        if technical:
            sections["technical"] = technical

    return _build_response(ext.upper().strip('.') or "AUDIO", sections)


def _extract_video_metadata(file_path):
    import cv2

    ext = os.path.splitext(file_path)[1].lower()
    sections = {"file_info": _file_info_section(file_path), "video": {}}

    try:
        cap = cv2.VideoCapture(file_path)
        fps = cap.get(cv2.CAP_PROP_FPS) or 0
        frame_count = cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0
        width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH) or 0)
        height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT) or 0)
        duration = round(frame_count / fps, 2) if fps else None
        cap.release()

        sections["video"] = {
            "Duration": duration,
            "Resolution": f"{width} x {height}" if width and height else None,
            "Frame Rate": round(fps, 2) if fps else None,
        }
    except Exception as e:
        sections["video"]["Note"] = str(e)

    hachoir_map = _collect_hachoir_map(file_path)
    if hachoir_map:
        technical = {}
        for key, value in hachoir_map.items():
            if any(term in key.lower() for term in ("date", "time", "device", "model", "gps", "location", "creator", "software", "duration", "resolution")):
                technical[key] = value
        if technical:
            sections["video"].update({k: v for k, v in technical.items() if k not in sections["video"]})
            sections["technical"] = technical

    return _build_response(ext.upper().strip('.') or "VIDEO", sections)


def _extract_generic_metadata(file_path):
    hachoir_map = _collect_hachoir_map(file_path)
    sections = {
        "file_info": _file_info_section(file_path),
        "technical": hachoir_map or {"Note": "No additional metadata found."},
    }
    return _build_response(os.path.splitext(file_path)[1].upper().strip('.') or "UNKNOWN", sections)


def _extract_jpeg_metadata(file_path):
    """Extract JPEG EXIF metadata using piexif."""
    import piexif

    try:
        exif_dict = piexif.load(file_path)
    except Exception:
        return {
            "file_type": os.path.splitext(file_path)[1].upper().strip('.'),
            "metadata": {
                "File Name": os.path.basename(file_path),
                "Note": "No EXIF metadata found."
            }
        }

    def decode_value(value):
        if isinstance(value, bytes):
            return value.decode("utf-8", errors="ignore").rstrip("\x00")
        return str(value)

    info = {
        "file_type": os.path.splitext(file_path)[1].upper().strip('.'),
        "metadata": {
            "File Name": os.path.basename(file_path),
        }
    }

    for field, (ifd, tag) in JPEG_EDITABLE_FIELDS.items():
        value = exif_dict.get(ifd, {}).get(tag)
        if value is not None:
            info["metadata"][field] = decode_value(value)

    if len(info["metadata"]) == 1:
        info["metadata"]["Note"] = "No EXIF metadata found."

    return info
